import io
from pathlib import Path
import pandas as pd
import plotly.graph_objects as go
import streamlit as st
from openpyxl import load_workbook

st.set_page_config(page_title='WEALTH OS | Financial Command Center', page_icon='◆', layout='wide', initial_sidebar_state='expanded')
st.markdown('''<style>
.stApp {background:#0B111B;color:#E8EDF6} [data-testid="stSidebar"]{background:#111C2B}
[data-testid="stMetric"]{background:#142235;border:1px solid #263C53;border-radius:15px;padding:17px}
[data-testid="stMetricLabel"]{color:#A6B8C9} h1,h2,h3{letter-spacing:-.03em}
.stButton>button,.stDownloadButton>button{background:#C6A46A;color:#101A26;border:0;border-radius:10px;font-weight:700}
</style>''',unsafe_allow_html=True)
BASE=Path(__file__).parent/'financial_goals_tracker_2028_2040.xlsx'
DEFAULT=BASE.read_bytes()
# The workbook is the source of truth. Map UI fields to actual Inputs row labels,
# rather than assuming the dashboard and Excel have identical row numbers.
FIELD_NAMES={
2:'Starting cash Jan 2028',
3:'Jan 2028 bonus',
4:'2028 pre-switch monthly net',
5:'New gross salary July 2028',
6:'Net factor',
7:'Annual salary raise',
8:'Promotion year',
9:'Promotion raise',
10:'Bonus in months',
11:'First annual bonus year',
12:'Living expenses pre-marriage',
13:'Living expenses after marriage',
14:'Marriage year',
15:'Marriage cost',
16:'Car 1 price',
17:'Car 1 down payment',
18:'Car 1 balloon percentage',
19:'Car 1 financing annual rate',
20:'Car 1 term in months',
21:'Car 1 purchase year',
22:'Car 1 sale year',
23:'Car 1 resale value',
24:'Car 2 purchase year',
25:'Car 2 cash price',
26:'Villa year',
27:'Villa price',
28:'Villa down payment',
29:'Villa fees and fitout',
30:'Mortgage rate',
31:'Mortgage years',
32:'Emergency reserve',
33:'Annual portfolio return',
34:'Portfolio seed after villa',
35:'Monthly investing after villa',
36:'Bonus invested after villa',
37:'Withdrawal per year'}

LABELS={i:name for i,name in FIELD_NAMES.items()}
PCT={6,7,9,18,19,30,33,36}
YEARS={8,11,14,20,21,22,24,26,31}

def workbook_rows(w):
    ws=w['Inputs']
    mapping={str(ws.cell(row,1).value).strip():row for row in range(2,ws.max_row+1)
             if ws.cell(row,1).value is not None}
    missing=set(FIELD_NAMES.values())-set(mapping)
    if missing:raise ValueError('Missing required Excel inputs: '+', '.join(sorted(missing)))
    return mapping

def read_inputs(data):
    w=load_workbook(io.BytesIO(data),read_only=True,data_only=False)
    mapping=workbook_rows(w)
    d={i:float(w['Inputs'].cell(mapping[name],2).value or 0) for i,name in FIELD_NAMES.items()}
    w.close()
    return d

def amort_payment(principal,annual,months,balloon=0):
 if months<=0:return 0
 r=annual/12
 return (principal-balloon/(1+r)**months)*r/(1-(1+r)**(-months)) if r else (principal-balloon)/months

def balance(principal,annual,payment,n):
 r=annual/12
 return principal*(1+r)**n-payment*((1+r)**n-1)/r if r else principal-payment*n

def simulate(d):
 p=d[16]-d[17]; balloon=d[16]*d[18]; term=int(d[20]); carpay=amort_payment(p,d[19],term,balloon)
 mortgage=amort_payment(max(0,d[27]-d[28]),d[30],int(d[31])*12)
 cash=d[2]; portfolio=0; rows=[]; lowest=1e18; monthly=[]
 purchase1=(int(d[21]),7); sale1=(int(d[22]),7); purchase2=(int(d[24]),7); marriage=(int(d[14]),1); villa=(int(d[26]),12)
 car_months=0; loan_balance=p; mortgage_started=False
 for year in range(2028,2041):
  income=spending=goalcost=bonus_total=investment_total=0
  for month in range(1,13):
   if year==2028: gross=d[5]; net=d[4] if month<=6 else gross*d[6]
   else: gross=d[5]*(1+d[7])**(year-2028)*(1+d[9] if year>=int(d[8]) else 1);net=gross*d[6]
   bonus=(d[3] if (year,month)==(2028,1) else gross*d[10]*d[6] if year>=int(d[11]) and month==12 else 0)
   living=d[12] if (year,month)<marriage else d[13]
   flow=net+bonus-living; income+=net+bonus; bonus_total+=bonus; spending+=living
   if (year,month)==purchase1:flow-=d[17];goalcost+=d[17]
   if purchase1<=(year,month)<sale1 and car_months<term:
    flow-=carpay; spending+=carpay; car_months+=1
   if (year,month)==sale1:
    payoff=max(0,balance(p,d[19],carpay,car_months)); flow+=d[23]-payoff;goalcost+=payoff-d[23]
   if (year,month)==purchase2:flow-=d[25];goalcost+=d[25]
   if (year,month)==marriage:flow-=d[15];goalcost+=d[15]
   if (year,month)==villa:
    flow-=d[28]+d[29];goalcost+=d[28]+d[29];mortgage_started=True
    seed=min(max(0,cash+flow-d[32]),d[34]);flow-=seed;portfolio+=seed;investment_total+=seed
   elif mortgage_started:
    flow-=mortgage;spending+=mortgage
    invest=min(max(0,cash+flow-d[32]),d[35]+(bonus*d[36] if bonus else 0))
    flow-=invest;portfolio+=invest;investment_total+=invest
   if mortgage_started:
    portfolio*=((1+d[33])**(1/12))
    if month==12 and year>int(d[26]):portfolio-=min(max(0,portfolio),d[37])
   cash+=flow;lowest=min(lowest,cash)
   monthly.append({'year':year,'month':month,'cash':cash,'portfolio':portfolio,'net_cash_flow':flow})
  rows.append({'السنة':year,'الدخل':income,'البونص':bonus_total,'المعيشة والأقساط':spending,'أحداث مالية صافية':goalcost,'استثمار جديد':investment_total,'السيولة':cash,'المحفظة':portfolio,'الثروة السائلة':cash+portfolio})
 return pd.DataFrame(rows),pd.DataFrame(monthly),carpay,mortgage,lowest

st.markdown('## ◆ WEALTH OS')
st.caption('PERSONAL FINANCIAL COMMAND CENTER  ·  2028—2040  ·  SAR')
upload=st.sidebar.file_uploader('📂 ارفع ملف الإكسل المحدّث',type=['xlsx'])
source=upload.getvalue() if upload else DEFAULT
try: values=read_inputs(source)
except Exception as e:st.error(f'تعذّر قراءة الملف: {e}');st.stop()
st.sidebar.markdown('### ⚙️ محاكي الافتراضات')
if 'last_file' not in st.session_state or st.session_state.last_file!=hash(source):
 st.session_state.last_file=hash(source)
 for i,v in values.items():st.session_state[f'input_{i}']=float(v)
with st.sidebar.expander('الدخل والبونص',expanded=True):
 for i in [5,6,7,8,9,10,11,4,3]:
  step=.01 if i in PCT else 1.0 if i in YEARS else 500.0
  st.number_input(LABELS[i],key=f'input_{i}',step=step,format='%.3f' if i in PCT else '%.0f')
with st.sidebar.expander('السيارات والزواج'):
 for i in [12,13,14,15,16,17,18,19,20,21,22,23,24,25]:
  st.number_input(LABELS[i],key=f'input_{i}',step=.01 if i in PCT else 1.0 if i in YEARS else 1000.0,format='%.3f' if i in PCT else '%.0f')
with st.sidebar.expander('الفيلا والمحفظة'):
 for i in [26,27,28,29,30,31,32,33,34,35,36,37,2]:
  st.number_input(LABELS[i],key=f'input_{i}',step=.01 if i in PCT else 1.0 if i in YEARS else 1000.0,format='%.3f' if i in PCT else '%.0f')
d={i:float(st.session_state[f'input_{i}']) for i in LABELS}
try:annual,monthly,carpay,mortgage,lowest=simulate(d)
except Exception as e:st.error(f'تحقق من الافتراضات: {e}');st.stop()

def sar(x):return f'{x:,.0f} ر.س'
selected=st.select_slider('السنة المعروضة',options=list(range(2028,2041)),value=2035)
r=annual[annual['السنة']==selected].iloc[0]
a,b,c,e=st.columns(4)
a.metric('الثروة السائلة المتوقعة',sar(r['الثروة السائلة']))
b.metric('السيولة النقدية',sar(r['السيولة']))
c.metric('المحفظة الاستثمارية',sar(r['المحفظة']))
e.metric('الدخل السنوي',sar(r['الدخل']))
if lowest<0:st.error(f'⚠️ الخطة تُظهر عجزًا نقديًا مؤقتًا يصل إلى {sar(-lowest)}. يلزم تعديل توقيت الأهداف أو التمويل.')
else:st.success('✓ لا يظهر عجز نقدي في التدفقات الشهرية وفق الافتراضات المدخلة. هذا لا يغني عن اختبار المخاطر.')
t1,t2,t3,t4=st.tabs(['📊 نظرة عامة','🎯 الأهداف','💸 التدفقات النقدية','🧪 اختبار السيناريوهات'])
with t1:
 fig=go.Figure()
 for col,color in [('السيولة','#4DB9B4'),('المحفظة','#D3AD71'),('الثروة السائلة','#8D9FF2')]:
  fig.add_trace(go.Scatter(x=annual['السنة'],y=annual[col],mode='lines+markers',name=col,line=dict(color=color,width=3)))
 fig.update_layout(template='plotly_dark',paper_bgcolor='#0B111B',plot_bgcolor='#0B111B',title='تطور السيولة والمحفظة والثروة السائلة',height=420,legend=dict(orientation='h'),yaxis_title='ريال')
 st.plotly_chart(fig,use_container_width=True)
 c1,c2=st.columns(2);c1.metric('قسط سيارة 2028 التقديري',sar(carpay));c2.metric('القسط العقاري التقديري',sar(mortgage))
 st.caption('الثروة السائلة = النقد + المحفظة فقط. لا تشمل القيمة السوقية للفيلا والسيارات أو رصيد التمويل العقاري.')
with t2:
 goals=[('🚘 السيارة الأولى',int(d[21]),d[16]),('💍 الزواج',int(d[14]),d[15]),('🚘 السيارة الثانية',int(d[24]),d[25]),('🏡 الفيلا',int(d[26]),d[27]),('📈 محفظة مليون',2038,1000000)]
 for title,yr,amt in goals:
  st.markdown(f'**{title}**  ·  {yr}  ·  {sar(amt)}')
  if title.startswith('📈'):
   actual=float(annual.loc[annual['السنة']==yr,'المحفظة'].iloc[0]);st.progress(min(1,max(0,actual/amt)));st.caption(f'المحفظة المتوقعة: {sar(actual)}')
  else:st.caption('موعد وتكلفة مخططان — تحقق من الرصيد المتاح وقت التنفيذ')
with t3:
 st.markdown('#### سجل الأرقام الفعلية من Excel — Update Log')
 try:
  histwb=load_workbook(io.BytesIO(source),read_only=True,data_only=True)
  histws=histwb['Update Log']
  hist=list(histws.iter_rows(min_row=4,max_col=7,values_only=True))
  hist=[r for r in hist if r[0] is not None]
  if hist:
   st.dataframe(pd.DataFrame(hist,columns=['التاريخ','صافي الراتب الفعلي','البونص الفعلي','الكاش الفعلي','الاستثمارات الفعلية','الديون','ملاحظات']),hide_index=True,use_container_width=True)
  else:st.info('أضف أرقامك الفعلية في تبويب Update Log في الإكسل ثم ارفعه.')
  histwb.close()
 except Exception as exc:st.warning(f'تعذر قراءة Update Log: {exc}')
 fig2=go.Figure()
 fig2.add_bar(x=annual['السنة'],y=annual['الدخل'],name='الدخل',marker_color='#4DB9B4')
 fig2.add_bar(x=annual['السنة'],y=-annual['المعيشة والأقساط'],name='المعيشة والأقساط',marker_color='#DB7B76')
 fig2.update_layout(template='plotly_dark',paper_bgcolor='#0B111B',plot_bgcolor='#0B111B',barmode='relative',height=370)
 st.plotly_chart(fig2,use_container_width=True)
 st.dataframe(annual.style.format({col:'{:,.0f}' for col in annual.columns if col!='السنة'}),hide_index=True,use_container_width=True)
with t4:
 st.markdown('#### أثر العائد الاستثماري على محفظتك')
 comparison=[]
 for rate in [.05,.08,.10]:
  scenario=dict(d);scenario[33]=rate;out,_,_,_,_=simulate(scenario)
  comparison.append({'العائد المفترض':f'{rate:.0%}','محفظة 2038':out.loc[out['السنة']==2038,'المحفظة'].iloc[0],'محفظة 2040':out.loc[out['السنة']==2040,'المحفظة'].iloc[0]})
 st.dataframe(pd.DataFrame(comparison).style.format({'محفظة 2038':'{:,.0f}','محفظة 2040':'{:,.0f}'}),hide_index=True,use_container_width=True)
 st.warning('العوائد افتراضية وليست مضمونة. المحاكاة لا تشمل التضخم، الرسوم، ضرائب محتملة، الأطفال، أو هبوط الأسواق المفاجئ.')

# Download the updated source workbook: input cells are synced, formulas preserved.
wb=load_workbook(io.BytesIO(source));ws=wb['Inputs']
for i,v in d.items():
    row=workbook_rows(wb)[FIELD_NAMES[i]]
    ws.cell(row,2).value=int(v) if i in YEARS else v
try:wb.calculation.fullCalcOnLoad=True;wb.calculation.forceFullCalc=True
except Exception:pass
buf=io.BytesIO();wb.save(buf)
st.sidebar.download_button('⬇️ تنزيل الإكسل بالافتراضات الجديدة',buf.getvalue(),file_name='wealth_plan_updated.xlsx',mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',use_container_width=True)
st.sidebar.caption('Excel Inputs ↔ Dashboard: القيم تُقرأ مباشرة من تبويب Inputs وتُكتب إلى نفس الخلايا عند التنزيل. لا يوجد حفظ تلقائي على خادم.')
st.caption('مرتبط بملف Excel: Inputs للتحكم، Update Log للأرقام الفعلية، وMonthly Plan / Annual Summary للتوقعات. قد توجد فروقات طفيفة في توقيت الأقساط؛ راجع شروط التمويل الفعلية.')
